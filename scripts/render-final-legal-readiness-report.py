from __future__ import annotations

import argparse
import html
import re
from dataclasses import dataclass
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "FINAL_LEGAL_READINESS_REPORT.md"
DOCX_OUT = ROOT / "FINAL_LEGAL_READINESS_REPORT.docx"
PDF_OUT = ROOT / "FINAL_LEGAL_READINESS_REPORT.pdf"


@dataclass
class Block:
    kind: str
    value: object


def normalize_display(text: str) -> str:
    return (
        text.replace("\u2010", "-")
        .replace("\u2011", "-")
        .replace("\u2012", "-")
        .replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2212", "-")
        .replace("\u2192", " -> ")
    )


def split_row(line: str) -> list[str]:
    value = line.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|"):
        value = value[:-1]
    cells = re.split(r"(?<!\\)\|", value)
    return [normalize_display(cell.replace("\\|", "|").strip()) for cell in cells]


def is_separator(line: str) -> bool:
    cells = split_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def parse_markdown(text: str) -> list[Block]:
    lines = text.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.strip().startswith("```"):
            fence = line.strip()
            code: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            blocks.append(Block("code", "\n".join(code)))
            continue
        if line.lstrip().startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            if len(table_lines) >= 2 and is_separator(table_lines[1]):
                headers = split_row(table_lines[0])
                rows: list[list[str]] = []
                for row_line in table_lines[2:]:
                    if is_separator(row_line):
                        continue
                    cells = split_row(row_line)
                    if len(cells) < len(headers):
                        cells.extend([""] * (len(headers) - len(cells)))
                    elif len(cells) > len(headers):
                        cells = cells[: len(headers) - 1] + [" | ".join(cells[len(headers) - 1 :])]
                    rows.append(cells)
                blocks.append(Block("table", (headers, rows)))
            else:
                blocks.append(Block("paragraph", " ".join(line.strip().strip("|") for line in table_lines)))
            continue
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            blocks.append(Block("heading", (len(heading.group(1)), heading.group(2).strip())))
            i += 1
            continue
        if re.match(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$", line):
            blocks.append(Block("rule", ""))
            i += 1
            continue
        if re.match(r"^\s*(?:[-+*]\s+|\d+[.)]\s+)", line):
            items: list[tuple[str, str]] = []
            while i < len(lines) and re.match(r"^\s*(?:[-+*]\s+|\d+[.)]\s+)", lines[i]):
                match = re.match(r"^\s*(?:([-+*])|(\d+[.)]))\s+(.*)$", lines[i])
                assert match is not None
                items.append(("bullet" if match.group(1) else "number", match.group(3).strip()))
                i += 1
            blocks.append(Block("list", items))
            continue
        if line.lstrip().startswith(">"):
            quote: list[str] = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            blocks.append(Block("quote", " ".join(quote)))
            continue
        paragraph = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(("#", "|", "```", ">")) and not re.match(r"^\s*(?:[-+*]\s+|\d+[.)]\s+|(?:-{3,}|\*{3,}|_{3,})\s*$)", lines[i]):
            paragraph.append(lines[i].strip())
            i += 1
        blocks.append(Block("paragraph", " ".join(paragraph)))
    return blocks


def clean_cell_text(text: str) -> str:
    value = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
    value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", value)
    value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
    value = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", value)
    value = value.replace("`", "")
    value = re.sub(r"\\([\\`*_{}\[\]()#+.!|>-])", r"\1", value)
    return normalize_display(value).strip()


INLINE_RE = re.compile(r"(\*\*.+?\*\*|__.+?__|`.+?`|(?<!\*)\*[^*]+\*(?!\*))")


def inline_parts(text: str) -> list[tuple[str, str]]:
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = normalize_display(text)
    parts: list[tuple[str, str]] = []
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith(("**", "__")) and token.endswith(("**", "__")):
            parts.append(("bold", token[2:-2]))
        elif token.startswith("`") and token.endswith("`"):
            parts.append(("code", token[1:-1]))
        elif token.startswith("*") and token.endswith("*"):
            parts.append(("italic", token[1:-1]))
        else:
            parts.append(("plain", token.replace("\\|", "|")))
    return parts


def table_groups(headers: list[str], rows: list[list[str]]) -> list[tuple[list[int], list[str], list[list[str]]]]:
    count = len(headers)
    if count <= 8:
        return [(list(range(count)), headers, rows)]
    if count > 14:
        first_end = (count + 2) // 3
        second_end = (2 * count + 2) // 3
        index_groups = [
            list(range(first_end)),
            [0] + list(range(first_end, second_end)),
            [0] + list(range(second_end, count)),
        ]
        return [
            (indices, [headers[index] for index in indices], [[row[index] for index in indices] for row in rows])
            for indices in index_groups
        ]
    split = (count + 1) // 2
    left = list(range(split))
    right = [0] + list(range(split, count))
    return [
        (left, [headers[index] for index in left], [[row[index] for index in left] for row in rows]),
        (right, [headers[index] for index in right], [[row[index] for index in right] for row in rows]),
    ]


def starts_wide_group(blocks: list[Block], index: int) -> bool:
    if index >= len(blocks):
        return False
    block = blocks[index]
    if block.kind == "table" and len(block.value[0]) > 6:
        return True
    return (
        block.kind == "heading"
        and index + 1 < len(blocks)
        and blocks[index + 1].kind == "table"
        and len(blocks[index + 1].value[0]) > 6
    )


def render_docx(blocks: list[Block]) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Emu, Inches, Pt, RGBColor

    document = Document()
    props = document.core_properties
    props.title = "Informe final de preparación legal"
    props.subject = "Resultado del generador jurídico y evidencia del Caso 01"
    props.author = ""
    props.keywords = "generador jurídico, validación, Caso 01"

    def set_section(section, landscape: bool) -> None:
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        if landscape:
            section.page_width, section.page_height = Inches(11), Inches(8.5)
            section.left_margin = section.right_margin = Inches(0.52)
            section.top_margin, section.bottom_margin = Inches(0.58), Inches(0.62)
        else:
            section.page_width, section.page_height = Inches(8.5), Inches(11)
            section.left_margin = section.right_margin = Inches(0.72)
            section.top_margin, section.bottom_margin = Inches(0.72), Inches(0.68)
        section.header_distance = Inches(0.28)
        section.footer_distance = Inches(0.3)

    set_section(document.sections[0], False)

    def font_for_style(style, size, color, bold=False):
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = bold
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    normal = document.styles["Normal"]
    font_for_style(normal, 9.6, "263746")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    font_for_style(document.styles["Title"], 22, "17324D", True)
    document.styles["Title"].paragraph_format.space_after = Pt(8)
    for name, size, color in [("Heading 1", 15, "17324D"), ("Heading 2", 12, "245578"), ("Heading 3", 10.5, "245578"), ("Heading 4", 9.8, "34495E")]:
        style = document.styles[name]
        font_for_style(style, size, color, True)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    for section in document.sections:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("INFORME TÉCNICO  |  GENERADOR JURÍDICO")
        run.font.name = "Arial"
        run.font.size = Pt(7.3)
        run.font.color.rgb = RGBColor.from_string("6B7785")
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run("Informe de preparación legal   |   Página ")
        run.font.name = "Arial"
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string("6B7785")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)

    def add_runs(paragraph, value: str) -> None:
        for kind, text in inline_parts(value):
            chunks = text.split("\n")
            for idx, chunk in enumerate(chunks):
                if idx:
                    paragraph.add_run().add_break()
                if not chunk:
                    continue
                run = paragraph.add_run(chunk)
                if kind == "bold":
                    run.bold = True
                elif kind == "italic":
                    run.italic = True
                elif kind == "code":
                    run.font.name = "Consolas"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
                    run.font.size = Pt(8.4)
                    run.font.color.rgb = RGBColor.from_string("245578")

    def set_cell_shading(cell, fill: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = tc_pr.find(qn("w:shd"))
        if shd is None:
            shd = OxmlElement("w:shd")
            tc_pr.append(shd)
        shd.set(qn("w:fill"), fill)

    def set_cell_margins(cell, top=55, start=65, bottom=55, end=65) -> None:
        tc = cell._tc
        tc_pr = tc.get_or_add_tcPr()
        margins = tc_pr.first_child_found_in("w:tcMar")
        if margins is None:
            margins = OxmlElement("w:tcMar")
            tc_pr.append(margins)
        for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
            node = margins.find(qn(f"w:{edge}"))
            if node is None:
                node = OxmlElement(f"w:{edge}")
                margins.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    def set_repeat_table_header(row) -> None:
        tr_pr = row._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)

    def set_no_split(row) -> None:
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))

    def add_docx_table(headers: list[str], rows: list[list[str]]) -> None:
        table = document.add_table(rows=1, cols=len(headers))
        table.autofit = False
        table.style = "Table Grid"
        section = document.sections[-1]
        usable = section.page_width - section.left_margin - section.right_margin
        sample = [headers] + rows[:12]
        weights = [max(5, min(26, max(len(clean_cell_text(row[c])) for row in sample))) for c in range(len(headers))]
        total_weight = sum(weights)
        widths = [usable * weight / total_weight for weight in weights]
        min_width = Inches(0.52)
        for idx, width in enumerate(widths):
            table.columns[idx].width = Emu(int(round(max(min_width, width))))
        for row_idx, values in enumerate([headers] + rows):
            cells = table.rows[0].cells if row_idx == 0 else table.add_row().cells
            set_no_split(table.rows[row_idx])
            for col, cell in enumerate(cells):
                cell.width = table.columns[col].width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                set_cell_margins(cell)
                if row_idx == 0:
                    set_cell_shading(cell, "17324D")
                elif row_idx % 2 == 0:
                    set_cell_shading(cell, "F1F5F8")
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                for line_idx, line in enumerate(clean_cell_text(values[col]).split("\n")):
                    if line_idx:
                        para.add_run().add_break()
                    run = para.add_run(line)
                    run.font.name = "Arial"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
                    run.font.size = Pt(7.0 if len(headers) >= 7 else 8.0)
                    run.bold = row_idx == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if row_idx == 0 else "263746")
        set_repeat_table_header(table.rows[0])
        document.add_paragraph().paragraph_format.space_after = Pt(1)

    def wide_table_at(index: int) -> bool:
        return blocks[index].kind == "table" and len(blocks[index].value[0]) > 6

    landscape = False
    for index, block in enumerate(blocks):
        if block.kind == "heading" and not landscape and index + 1 < len(blocks) and wide_table_at(index + 1):
            section = document.add_section(WD_SECTION.NEW_PAGE)
            set_section(section, True)
            landscape = True
        elif block.kind == "table" and len(block.value[0]) > 6 and not landscape:
            section = document.add_section(WD_SECTION.NEW_PAGE)
            set_section(section, True)
            landscape = True

        if block.kind == "heading":
            level, value = block.value
            para = document.add_paragraph(style="Title" if level == 1 else f"Heading {min(level, 4)}")
            add_runs(para, value)
        elif block.kind == "paragraph":
            para = document.add_paragraph()
            add_runs(para, block.value)
        elif block.kind == "quote":
            para = document.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.28)
            para.paragraph_format.right_indent = Inches(0.18)
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(7)
            add_runs(para, block.value)
            for run in para.runs:
                run.italic = True
                run.font.color.rgb = RGBColor.from_string("4D6071")
        elif block.kind == "code":
            para = document.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.15)
            para.paragraph_format.space_after = Pt(5)
            para.paragraph_format.line_spacing = 1.0
            for line_idx, line in enumerate(block.value.split("\n")):
                if line_idx:
                    para.add_run().add_break()
                run = para.add_run(normalize_display(line))
                run.font.name = "Consolas"
                run.font.size = Pt(7.6)
                run.font.color.rgb = RGBColor.from_string("364B5E")
        elif block.kind == "list":
            for list_kind, value in block.value:
                style = "List Bullet" if list_kind == "bullet" else "List Number"
                para = document.add_paragraph(style=style)
                para.paragraph_format.space_after = Pt(2)
                add_runs(para, value)
        elif block.kind == "table":
            headers, rows = block.value
            for _, group_headers, group_rows in table_groups(headers, rows):
                add_docx_table(group_headers, group_rows)
            if landscape and not starts_wide_group(blocks, index + 1):
                section = document.add_section(WD_SECTION.NEW_PAGE)
                set_section(section, False)
                landscape = False
        elif block.kind == "rule":
            document.add_paragraph().paragraph_format.space_after = Pt(3)

    document.save(DOCX_OUT)


def render_pdf(blocks: list[Block]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        BaseDocTemplate,
        Frame,
        KeepTogether,
        LongTable,
        NextPageTemplate,
        PageBreak,
        PageTemplate,
        Paragraph,
        Spacer,
        TableStyle,
    )

    navy = colors.HexColor("#17324D")
    blue = colors.HexColor("#245578")
    ink = colors.HexColor("#263746")
    muted = colors.HexColor("#687887")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=21, leading=25, textColor=navy, spaceAfter=11, alignment=TA_LEFT, keepWithNext=True))
    for name, size, leading, color in [("ReportH1", 15, 18, navy), ("ReportH2", 12, 15, blue), ("ReportH3", 10.2, 13, blue), ("ReportH4", 9.4, 12, ink)]:
        styles.add(ParagraphStyle(name=name, parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=size, leading=leading, textColor=color, spaceBefore=10, spaceAfter=4, keepWithNext=True))
    styles.add(ParagraphStyle(name="ReportBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12, textColor=ink, spaceAfter=5, splitLongWords=1))
    styles.add(ParagraphStyle(name="ReportQuote", parent=styles["ReportBody"], leftIndent=18, rightIndent=12, textColor=muted, fontName="Helvetica-Oblique", spaceBefore=3, spaceAfter=7))
    styles.add(ParagraphStyle(name="ReportCode", parent=styles["Code"], fontName="Courier", fontSize=7.2, leading=9, textColor=ink, leftIndent=9, rightIndent=5, spaceBefore=2, spaceAfter=6, splitLongWords=1))
    styles.add(ParagraphStyle(name="ReportTable", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.3, leading=9, textColor=ink, spaceAfter=0, splitLongWords=1, wordWrap="CJK"))
    styles.add(ParagraphStyle(name="ReportTableSmall", parent=styles["ReportTable"], fontSize=6.5, leading=8))
    styles.add(ParagraphStyle(name="ReportTableHeader", parent=styles["ReportTable"], fontName="Helvetica-Bold", fontSize=7.1, leading=8.4, textColor=navy))
    styles.add(ParagraphStyle(name="ReportBullet", parent=styles["ReportBody"], leftIndent=17, firstLineIndent=-10, spaceAfter=2))

    class ReportDoc(BaseDocTemplate):
        def __init__(self, filename: str):
            super().__init__(filename, title="Informe final de preparación legal", author="", subject="Resultado del generador jurídico y evidencia del Caso 01")
            portrait_frame = Frame(0.72 * inch, 0.62 * inch, 7.06 * inch, 9.65 * inch, id="portrait", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
            landscape_frame = Frame(0.52 * inch, 0.62 * inch, 10.0 * inch, 7.22 * inch, id="landscape", leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
            self.addPageTemplates([
                PageTemplate(id="portrait", pagesize=letter, frames=portrait_frame, onPage=self.decorate),
                PageTemplate(id="landscape", pagesize=landscape(letter), frames=landscape_frame, onPage=self.decorate),
            ])

        def decorate(self, canvas, doc):
            width, height = doc.pagesize
            canvas.saveState()
            canvas.setStrokeColor(colors.HexColor("#D6DEE5"))
            canvas.setLineWidth(0.45)
            canvas.line(0.72 * inch if width < height else 0.52 * inch, 0.48 * inch, width - (0.72 * inch if width < height else 0.52 * inch), 0.48 * inch)
            canvas.setFont("Helvetica", 7.2)
            canvas.setFillColor(muted)
            canvas.drawString(0.72 * inch if width < height else 0.52 * inch, 0.31 * inch, "Informe de preparación legal")
            canvas.drawRightString(width - (0.72 * inch if width < height else 0.52 * inch), 0.31 * inch, f"Página {doc.page}")
            canvas.restoreState()

    def xml_inline(text: str) -> str:
        result: list[str] = []
        for kind, value in inline_parts(text):
            escaped = html.escape(value, quote=False).replace("\n", "<br/>")
            if kind == "bold":
                escaped = f"<b>{escaped}</b>"
            elif kind == "italic":
                escaped = f"<i>{escaped}</i>"
            elif kind == "code":
                escaped = f'<font name="Courier" color="#245578">{escaped}</font>'
            result.append(escaped)
        return "".join(result)

    def make_pdf_table(headers: list[str], rows: list[list[str]], page_width: float):
        parts = []
        for grouped_headers, grouped_rows in [(h, r) for _, h, r in table_groups(headers, rows)]:
            ncols = len(grouped_headers)
            style = styles["ReportTableSmall"] if ncols >= 7 else styles["ReportTable"]
            raw = [[Paragraph(xml_inline(cell), styles["ReportTableHeader"]) for cell in grouped_headers]]
            raw.extend([[Paragraph(xml_inline(cell), style) for cell in row] for row in grouped_rows])
            table_width = page_width
            sample = [grouped_headers] + grouped_rows[:12]
            weights = [max(5, min(26, max(len(clean_cell_text(row[c])) for row in sample))) for c in range(ncols)]
            total = sum(weights)
            col_widths = [table_width * value / total for value in weights]
            table = LongTable(raw, colWidths=col_widths, repeatRows=1, hAlign="LEFT", splitByRow=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE7F0")),
                ("TEXTCOLOR", (0, 0), (-1, 0), navy),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C9D3DC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F8")]),
            ]))
            parts.append(table)
            parts.append(Spacer(1, 5))
        return parts

    doc = ReportDoc(str(PDF_OUT))
    story = []
    landscape_active = False
    for index, block in enumerate(blocks):
        wide = block.kind == "table" and len(block.value[0]) > 6
        if block.kind == "heading" and not landscape_active and index + 1 < len(blocks) and blocks[index + 1].kind == "table" and len(blocks[index + 1].value[0]) > 6:
            story.extend([NextPageTemplate("landscape"), PageBreak()])
            landscape_active = True
        elif wide and not landscape_active:
            story.extend([NextPageTemplate("landscape"), PageBreak()])
            landscape_active = True

        if block.kind == "heading":
            level, value = block.value
            style_name = "ReportTitle" if level == 1 else f"ReportH{min(level, 4)}"
            story.append(Paragraph(xml_inline(value), styles[style_name]))
        elif block.kind == "paragraph":
            story.append(Paragraph(xml_inline(block.value), styles["ReportBody"]))
        elif block.kind == "quote":
            story.append(Paragraph(xml_inline(block.value), styles["ReportQuote"]))
        elif block.kind == "code":
            code = "<br/>".join(html.escape(normalize_display(line), quote=False).replace(" ", "&nbsp;") for line in block.value.split("\n"))
            story.append(Paragraph(code, styles["ReportCode"]))
        elif block.kind == "list":
            for list_kind, value in block.value:
                marker = "-" if list_kind == "bullet" else ""
                story.append(Paragraph(f"{marker} {xml_inline(value)}" if marker else xml_inline(value), styles["ReportBullet"]))
        elif block.kind == "table":
            page_width = 10.0 * inch if landscape_active else 7.06 * inch
            story.extend(make_pdf_table(*block.value, page_width))
            if landscape_active and not starts_wide_group(blocks, index + 1):
                story.extend([NextPageTemplate("portrait"), PageBreak()])
                landscape_active = False
        elif block.kind == "rule":
            story.append(Spacer(1, 4))

    doc.build(story)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", choices=("docx", "pdf"), required=True)
    args = parser.parse_args()
    source_text = SOURCE.read_text(encoding="utf-8")
    blocks = parse_markdown(source_text)
    table_count = sum(block.kind == "table" for block in blocks)
    coverage_rows = sum(len(block.value[1]) for block in blocks if block.kind == "table" and "Coverage ID" in block.value[0])
    print(f"Parsed {len(blocks)} blocks; {table_count} tables; {coverage_rows} rows in the Coverage table.")
    if args.format == "docx":
        render_docx(blocks)
        print(f"DOCX written: {DOCX_OUT}")
    else:
        render_pdf(blocks)
        print(f"PDF written: {PDF_OUT}")


if __name__ == "__main__":
    main()
