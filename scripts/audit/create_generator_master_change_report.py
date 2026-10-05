from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
AUDIT = ROOT / "audit" / "generator-master"
OUTPUT = AUDIT / "GENERATOR_MASTER_CHANGE_REPORT.docx"
MATRIX = json.loads((AUDIT / "writing-coverage.json").read_text(encoding="utf-8"))
SUMMARY = MATRIX["summary"]

NAVY = "17324D"
PALE_BLUE = "EAF1F7"
PALE_GRAY = "F3F5F7"
GRID = "D9DDE2"
BLACK = RGBColor(0, 0, 0)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    properties = table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), GRID)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_run_font(run, size=9.5, bold=None, color=BLACK, name="Arial") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def add_body(doc, text: str, *, bold_lead: str | None = None, italic=False):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        run = paragraph.add_run(text[len(bold_lead):])
    else:
        run = paragraph.add_run(text)
    set_run_font(run)
    run.italic = italic
    return paragraph


def add_bullet(doc, text: str):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.13)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(run, size=9.2)
    return paragraph


def add_heading(doc, text: str, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_run_font(run, size=13 if level == 1 else 11, bold=True)
    return paragraph


def add_table(doc, headers, rows, widths, *, font_size=8.1):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_borders(table)
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, title in enumerate(headers):
        cell = header.cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=125, bottom=125)
        set_cell_shading(cell, NAVY)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(title))
        set_run_font(run, size=font_size, bold=True, color=RGBColor(255, 255, 255))

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index, value in enumerate(values):
            cell = cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, PALE_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.04
            run = paragraph.add_run(str(value))
            set_run_font(run, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = paragraph.add_run("LEX PLANTILLAS  |  Generador Master  |  ")
    set_run_font(label, size=8, color=RGBColor(80, 90, 100))
    field_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    displayed = OxmlElement("w:t")
    displayed.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field_run._r.extend((begin, instruction, separate, displayed, end))
    set_run_font(field_run, size=8, color=RGBColor(80, 90, 100))


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.footer_distance = Inches(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = BLACK
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = BLACK
    doc.styles["Title"].font.size = Pt(23)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.size = Pt(13)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.size = Pt(11)
    doc.styles["Heading 2"].font.bold = True
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run("Reporte de cambios del generador jurídico")
    set_run_font(title_run, size=23, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run("LEX PLANTILLAS  |  Generador Master  |  4 de octubre de 2026")
    set_run_font(subtitle_run, size=10, color=RGBColor(70, 80, 90))

    add_body(
        doc,
        "Conclusión: se corrigieron errores concretos de clasificación y selección de tipos documentales, y se añadieron contratos estructurales parametrizados para los 214 tipos que el catálogo declaraba IMPLEMENTED. La regresión dirigida pasó (14 archivos, 310 tests), el typecheck pasó y el build de producción terminó. El cierre funcional global sigue incompleto: la matriz marca 0 de 297 filas PASS porque todavía no hay evidencia end-to-end de los 36 criterios por tipo."
    )

    add_heading(doc, "1. Alcance y resultado")
    add_body(doc, "Este reporte registra los cambios del checkpoint, los tests ejecutados, la auditoría de cuatro repositorios externos y el estado del generador. Se distingue el trabajo de este checkpoint de modificaciones que ya existían en el working tree para no atribuir cambios ajenos a esta fase.")
    metrics = [
        ("Tipos documentales canónicos", SUMMARY["canonicalDocumentCount"]),
        ("Identificadores de catálogo", SUMMARY["allCatalogIdentifiers"]),
        ("IDs legacy del selector histórico", SUMMARY["legacySelectableIds"]),
        ("Plantillas", SUMMARY["documentTemplateCount"]),
        ("Canónicos marcados IMPLEMENTED al inicio", SUMMARY["currentImplementedCount"]),
        ("Filas de la matriz deduplicada por ID", SUMMARY["rows"]),
        ("PASS de contrato completo de 36 criterios", SUMMARY["pass"]),
        ("FAIL conservadores", SUMMARY["fail"]),
        ("Bloqueados por dependencia externa", SUMMARY["blockedExternal"]),
        ("Opciones generables en Universal", len([r for r in MATRIX["rows"] if any("Universal" in s for s in r["uiSurface"])])),
        ("Opciones únicas seleccionables en las superficies medidas", len([r for r in MATRIX["rows"] if r["selectable"]])),
    ]
    add_table(doc, ["Medida", "Resultado"], metrics, [5.9, 1.25], font_size=8.6)
    add_body(doc, "FAIL significa que el contrato completo no está probado o que existe una carencia estructural; no afirma que cada fila sea un fallo runtime reproducido. La suite nueva cubre estructura y routing, no las 36 condiciones de generación, revisión y exportación.")

    add_heading(doc, "2. Cambios de producto")
    changes = [
        ("Catálogo", "La clasificación de formulario oficial y borrador asistido ahora tiene precedencia sobre la mera presencia de una plantilla. Una plantilla por sí sola ya no convierte esos tipos a IMPLEMENTED."),
        ("Routing", "El routing explícito rechaza un tipo cuyo estado canónico no es IMPLEMENTED, aunque exista DocumentTemplate; devuelve DOCUMENT_TYPE_NOT_IMPLEMENTED."),
        ("Contestaciones", "Se quitó la inferencia por regex del ID o etiqueta. El selector consume la capacidad explícita del catálogo: 13 outputs y un modo de redacción libre."),
        ("Universal", "Se creó una proyección legacy que solo entrega documentos implementados o alias activos con destino implementado y plantilla. Quince IDs de familia genérica dejaron de ser tipos de salida seleccionables; no se mapearon a escritos concretos por conjetura."),
        ("Contratos generados", "La suite paramétrica comprueba unicidad de ID y, para los 214 IMPLEMENTED, fuente compatible declarada, template, estructura, campos requeridos, estrategia/ruta coincidente y ausencia de fallback."),
        ("Inventario", "El builder ahora registra la suite estructural generada por cada tipo IMPLEMENTED y calcula conteos por materia y por las 107 familias documentales."),
    ]
    add_table(doc, ["Componente", "Cambio comprobable"], changes, [1.35, 5.8], font_size=8.25)

    add_heading(doc, "Archivos de producto y evidencia creados o ajustados", level=2)
    file_rows = [
        ("lib/catalog/legalCatalog.ts", "Precedencia de estados, capabilities del selector y proyecciones filtradas."),
        ("lib/legal-engine/documentRouting.ts", "Bloqueo de rutas para tipos canónicos que no están IMPLEMENTED."),
        ("app/machotes/components/CaseDocumentsReader.tsx", "Opciones de Contestaciones desde capability explícita."),
        ("app/machotes/page.tsx", "Selector Universal limitado a opciones generables del catálogo."),
        ("tests/legal-taxonomy/legalCatalogRegistry.test.ts", "RED/GREEN de estados, familias genéricas y selector Universal."),
        ("tests/legal-taxonomy/contestacionesCapability.test.ts; tests/components/contestacionesCatalogCapabilities.test.tsx", "Contrato de opciones de Contestaciones y verificación DOM."),
        ("tests/legal-taxonomy/implementedDocumentContracts.test.ts", "215 assertions estructurales generadas para unicidad y 214 tipos IMPLEMENTED."),
        ("scripts/audit/build-writing-coverage.mjs; tests/audit/generatorMasterInventory.test.ts", "Inventario automatizado, evidencia por familia y pruebas del propio reporte."),
        ("audit/generator-master/WRITING_COVERAGE_MATRIX.md; writing-coverage.json", "Matriz reproducible con los 297 IDs, estado, materia, familia, ruta y límites de evidencia."),
        ("audit/generator-master/STATUS.md; FAILURES.md; EXTERNAL_REPOS.md; WINDOWS_READINESS.md", "Checkpoint, causas, decisiones de reutilización y estado Windows."),
        ("scripts/audit/create_generator_master_change_report.py; GENERATOR_MASTER_CHANGE_REPORT.docx", "Builder reproducible y documento final de cambios."),
    ]
    add_table(doc, ["Archivo", "Contenido"], file_rows, [3.25, 3.9], font_size=7.7)

    add_heading(doc, "3. Ciclos RED y GREEN")
    cycles = [
        ("Formulario oficial marcado IMPLEMENTED", "18 pruebas: 16 pasaban y 2 fallaban. El ID de formulario oficial con template aparecía como IMPLEMENTED y el routing lo aceptaba.", "Se dio precedencia al estado REQUIRES_OFFICIAL_FORM y el routing empezó a validar el estado canónico.", "18/18 en legalCatalogRegistry.test.ts."),
        ("Contestaciones inferidas por nombre", "La prueba nueva falló: getContestacionesDocumentOptions no existía; el selector juntaba outputs con una regex.", "Capability contractual explícita en el catálogo; eliminada la mezcla dinámica por texto.", "Suite catálogo/UI dirigida PASS; 13 tipos + redacción libre confirmados."),
        ("Familias legacy presentadas como escritos", "19 pruebas: 18 pasaban y una fallaba porque el helper Universal no existía; 15 familias genéricas tenían ruta no generable.", "El selector usa solo tipos documentales IMPLEMENTED y alias activos con template; familias se conservan para navegación.", "21/21 en el paso catalog/inventory; 15 familias genéricas fuera de selección."),
        ("Cobertura por familia omitida", "Prueba de inventario falló: la suma de conteos por familia era 0 frente a 297 filas.", "El builder agrega el agregado byFamily y lo refleja en Markdown y JSON.", "217/217 entre inventario y contratos estructurales."),
    ]
    add_table(doc, ["Problema", "RED observado", "Corrección mínima", "GREEN"], cycles, [1.25, 2.15, 2.35, 1.4], font_size=7.35)

    add_heading(doc, "4. Regresión y compilación")
    add_table(doc, ["Verificación", "Resultado", "Alcance"], [
        ("Regresión dirigida", "14 archivos; 310/310 PASS", "Routing, matriz universal, catálogo, Contestaciones, contratos generados y regresiones de apelación cercanas."),
        ("Suite estructural generada", "214 documentos + unicidad PASS", "Solo invariantes estructurales; no equivale a generación jurídica ni exportación."),
        ("Typecheck", "PASS, exit code 0", "npm run typecheck después de los cambios de producto."),
        ("Build", "PASS, exit code 0", "npm run build; Prisma Client 6.19.3, Next 16.3.4 production webpack build y prepare-standalone.mjs: STANDALONE_PREPARED."),
    ], [1.55, 1.45, 4.15], font_size=8)
    add_body(doc, "No se llamó ningún provider real. Las pruebas de generación dentro de la regresión usan fixtures/stubs offline. No se ejecutó la suite contractual total, una generación real por familia, un E2E de editor/exportación, ni la verificación de archivos DOCX/PDF producidos para cada tipo.")

    add_heading(doc, "5. Estado por materia y familia")
    family_entries = sorted(SUMMARY["byFamily"].items(), key=lambda pair: (-pair[1]["count"], pair[0]))
    prominent = family_entries[:8]
    add_body(doc, f"La matriz registra {len(family_entries)} familias y 297 IDs. Todas tienen 0 PASS bajo el contrato completo de 36 criterios; la cifra expresa falta de evidencia end-to-end por tipo, no que cada documento haya fallado en ejecución. La lista completa está en WRITING_COVERAGE_MATRIX.md y writing-coverage.json.")
    add_table(doc, ["Familia", "IDs", "PASS", "FAIL"], [
        (family, value["count"], value["pass"], value["fail"]) for family, value in prominent
    ], [3.2, 1.1, 1.1, 1.1], font_size=8.2)
    add_body(doc, "Familias con más filas: " + "; ".join(f"{family} ({value['count']})" for family, value in prominent) + ". En todas las materias y familias permanecen pendientes las comprobaciones de contenido especializado, gates, DRAFT y exportación.")

    add_heading(doc, "6. Auditoría de repositorios externos")
    external_rows = [
        ("Magic UI", "d7207e5", "MIT", "Registry de componentes fuente y dependencias declaradas. Se evaluó para feedback/animación; no aporta al bloqueo del generador. No copiado."),
        ("1Code", "9f1bc76", "Apache-2.0; archivado 7 jul 2026", "Electron main/preload/renderer, IPC y gestión de procesos/worktrees. No se añadió una segunda app ni dependencias."),
        ("itsfree.ai", "c53a15a", "Sin LICENSE ni campo license en package.json", "Catálogo estático de proveedores/modelos; no es gateway ni política de privacidad. No se copió código/datos ni se añadió provider."),
        ("shadcn/ui", "295a1f1", "MIT", "CLI/registry resuelve fuente, destinos y dependencias. Patrón de registry útil, pero no se requiere migración de UI en esta fase."),
    ]
    add_table(doc, ["Repo", "Commit", "Licencia / estado", "Código inspeccionado y decisión"], external_rows, [0.9, 0.8, 1.65, 3.15], font_size=7.25)
    add_body(doc, "Los clones shallow/sparse quedaron fuera de APP-plantillas, en una carpeta temporal. Código externo copiado: ninguno. Dependencias añadidas: ninguna. Los commits y URLs se conservan en audit/generator-master/EXTERNAL_REPOS.md.")

    add_heading(doc, "7. Cambios previos del working tree preservados")
    add_body(doc, "Antes de este checkpoint ya había archivos modificados y pruebas nuevas sin seguimiento relacionados con la Fase 2b, apelaciones, el Manual Operativo y privacidad. No se revirtieron ni se atribuyen a los cambios del generador descritos arriba.")
    prior_groups = [
        "API y providers: app/api/legal-engine/appeal/reasoning-classification/route.ts; lib/ai/providers/gemini.ts, groq.ts y nvidia.ts.",
        "Motor y manual: lib/legal-engine/case-extraction/appealReasoningCandidates.ts, generatedLegalAdmission.ts, pendingFields.ts, qualityGate.ts y lib/operational-manual/core.ts.",
        "Pruebas ya existentes en el working tree: tests/audit/realAppealReasoningSource.test.ts; tests/legal-engine/appealReasoningImpact.test.ts y penalFallbackContract.test.ts; pruebas de importación, core y pureza del contexto del Manual Operativo.",
        "Pruebas nuevas ya presentes: tests/ai/providerErrorBodyPrivacy.test.ts; tests/api/appealReasoningClassificationAccess.test.ts; tests/legal-engine/pendingMarkerNormalization.test.ts.",
    ]
    for item in prior_groups:
        add_bullet(doc, item)

    add_heading(doc, "8. Bloqueos restantes y decisión")
    blockers = [
        "El catálogo aún marca 214 documentos como IMPLEMENTED sin evidencia end-to-end de los 36 criterios. La presencia de template y el nuevo test estructural no certifican contenido generado.",
        "No existe exportación real DRAFT DOCX/PDF comprobada por tipo ni E2E de las familias soportadas en este checkpoint.",
        "El conocido fallbackE2E de Amparo sigue pendiente de diagnóstico; no se cambió y no se considera exento del cierre final.",
        "Windows: Next build y preparación standalone pasaron, pero no se compiló ni probó un host/instalador de escritorio, persistencia local, offline, shutdown, firma o actualización.",
        "La aplicación no puede declararse GENERATOR_READY ni 100% funcional. La matriz completa conserva 0/297 PASS.",
    ]
    for item in blockers:
        add_bullet(doc, item)
    add_body(doc, "Siguiente trabajo recomendado: cerrar una familia a la vez con fixture anonimizado, prueba contractual generada desde el catálogo, generación offline/replay, gates intactos y exportación DRAFT real. Solo entonces elevar el estado de esa familia. Iniciar el empaquetado Windows después del cierre de las familias que la app siga ofreciendo.")

    add_heading(doc, "9. Verificación de este entregable")
    add_body(doc, "El archivo se reabrió con python-docx y la integridad CRC del paquete ZIP pasó; se comprobaron la presencia de resultados, límites y secciones del reporte. La inspección visual página por página quedó NO VERIFICADA: el renderer oficial no encontró LibreOffice (soffice.exe), y la exportación local con Microsoft Word COM falló con 0x80070520 por la sesión de Windows. No se instala software ni se envía este documento a un conversor externo.")

    doc.core_properties.title = "Reporte de cambios del generador jurídico"
    doc.core_properties.subject = "LEX PLANTILLAS Generador Master, pruebas y bloqueos"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "LEX PLANTILLAS, generador, auditoría, pruebas"
    doc.save(OUTPUT)
    print(f"DOCX_CREATED {OUTPUT}")


if __name__ == "__main__":
    build()
