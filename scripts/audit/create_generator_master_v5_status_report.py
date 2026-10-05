from __future__ import annotations

import importlib.util
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
AUDIT = ROOT / "audit" / "generator-master"
OUTPUT = AUDIT / "GENERATOR_MASTER_V5_STATUS_REPORT.docx"
SUMMARY = json.loads((AUDIT / "V5_EVIDENCE_SUMMARY.json").read_text(encoding="utf-8"))
GLOBAL = json.loads((AUDIT / "v5-global-current.json").read_text(encoding="utf-8"))
REPLAY = json.loads((AUDIT / "v5-replay" / "results.json").read_text(encoding="utf-8"))["records"]

STYLE_PATH = ROOT / "scripts" / "audit" / "create_generator_master_change_report.py"
spec = importlib.util.spec_from_file_location("generator_report_style", STYLE_PATH)
style = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(style)


def add_section(doc, title: str, intro: str | None = None) -> None:
    style.add_heading(doc, title, 1)
    if intro:
        style.add_body(doc, intro)


def add_path_bullets(doc, title: str, paths: list[str]) -> None:
    style.add_heading(doc, title, 2)
    for item in paths:
        style.add_bullet(doc, item)


def build() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    sec.footer_distance = Inches(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(style.qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(style.qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Title"].font.bold = True
    style.add_page_number(sec.footer.paragraphs[0])

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("Estado V5 del generador juridico")
    style.set_run_font(run, size=22, bold=True)
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    subrun = sub.add_run("LEX PLANTILLAS | Informe tecnico y funcional | 4 de octubre de 2026")
    style.set_run_font(subrun, size=10, color=RGBColor(65, 78, 91))

    style.add_body(
        doc,
        "Conclusion: V5 permanece BLOQUEADO. El replay offline recorrio 214 tipos y produjo DOCX/PDF estructuralmente validos, pero 0 de 277 tipos canonicos estan certificados funcionalmente. La regresion global actual conserva 25 fallas; npm run build no cerro por un bloqueo EPERM de la DLL de Prisma y el E2E real por interfaz no pudo abrir localhost:3200. No se llamaron proveedores reales.",
        bold_lead="Conclusion:",
    )

    add_section(doc, "1 Estado actual y metricas")
    metrics = [
        ("Tipos canonicos", "277"),
        ("Filas de la matriz", "297; 0 PASS"),
        ("Tipos con implementacion estructural", "214"),
        ("Replay offline / export estructural", "214/214 DOCX y 214/214 PDF"),
        ("Tipos certificados funcionalmente", "0/277"),
        ("FINAL permitido en replay", "0/214"),
        ("Regresion global actual", "4,053/4,123 PASS; 25 FAIL; 45 skipped"),
        ("Typecheck", "PASS; tsc --noEmit"),
        ("npm run build", "FAIL/BLOQUEADO; EPERM de DLL Prisma"),
        ("E2E Playwright por UI real", "BLOQUEADO antes de page load: EACCES localhost:3200"),
        ("Proveedores reales", "No llamados"),
        ("Revision humana", "Pendiente"),
    ]
    style.add_table(doc, ["Medida", "Resultado comprobado"], metrics, [3.1, 3.8], font_size=8.7)
    style.add_body(doc, "El subscore aislado sin hallazgos fue verdadero en 2 registros penales; no equivale al contrato completo, no promueve funcionalStatus y no permite FINAL. Los textos de fixtures/determinismo no prueban calidad juridica profesional.")

    add_section(doc, "2 Cambios principales")
    changes = [
        ("Demanda civil", "Los petitorios dependen de solicitudes confirmadas; no se agregan remedios favorables genericos. La personalidad solo se usa si esta confirmada; si falta, queda un pendiente especifico."),
        ("Contestaciones", "Las semillas legacy ya no afirman personalidad, domicilio, pruebas ni oportunidad sin soporte en el expediente."),
        ("Contrato de solicitud", "Una frase que prohibe un remedio no lo autoriza por coincidencia de palabras."),
        ("Pendientes y scorecard", "Los marcadores normalizados siguen representando dependencias abiertas. No se relajo ningun gate."),
        ("Estructura laboral", "El contrato incluye DERECHO cuando el template vigente la produce."),
        ("Apelacion", "Razonamientos por bloque con cita literal validable; la ruta mantiene auth, limites y consentimiento. No se probo provider real."),
        ("Catalogo y UI", "La disponibilidad depende de capacidad funcional explicita, no de inferencia por nombre. Los tipos no acreditados permanecen deshabilitados."),
        ("Privacidad y manual", "Los adapters omiten cuerpos de error en logs; el retrieval del manual distingue redaccion, perfil interno y formato de respuesta."),
        ("Evidencia", "Se incorporaron replay/scorecard por familia y un agregador que separa la corrida global exacta del agregado latest-per-file."),
    ]
    style.add_table(doc, ["Area", "Cambio"], changes, [1.55, 5.35], font_size=8.05)

    add_section(doc, "3 Pruebas RED y GREEN", "Los focales demuestran correcciones acotadas. No sustituyen la regresion global ni la revision profesional.")
    cycles = [
        ("Contratos base de familias", "958/970; 12 fallos iniciales", "335/335 focales"),
        ("Petitorios civiles y personalidad", "2 fallos reproducidos", "14/14 en cuatro archivos"),
        ("Scorecard", "4/4 fallan en fixture RED", "4/4"),
        ("Metodologia y scorecard", "1/5 falla", "15/15"),
        ("Marcadores/dependencias", "4 fallos de 7", "10/10"),
        ("Semillas legacy", "2 fallos de 10", "104/104"),
        ("Consentimiento judge", "2 fallos iniciales", "14/14 focales"),
        ("Apelacion UI", "3 fallos intermedios de 17", "17/17 focales"),
        ("Agregacion de evidencia", "No distingue corrida exacta", "4053/4123 validados contra JSON global"),
    ]
    style.add_table(doc, ["Ciclo", "RED observado", "GREEN focal"], cycles, [2.0, 2.45, 2.45], font_size=7.8)

    add_section(doc, "4 Replay por familia")
    fam = {}
    for record in REPLAY:
        fam.setdefault(record["family"], []).append(record)
    family_rows = []
    for name, records in sorted(fam.items()):
        mech = sum(r.get("mechanicalReplay") == "PASS" for r in records)
        score = sum(r.get("scorecard", {}).get("substantivePass") is True for r in records)
        final = sum(r.get("finalAllowed") is True for r in records)
        family_rows.append((name, len(records), f"{mech}/{len(records)}", score, final))
    family_rows.append(("TOTAL", len(REPLAY), "214/214", 2, 0))
    style.add_table(doc, ["Familia", "Casos", "Estructural", "Subscore", "FINAL"], family_rows,
                    [2.15, 0.75, 1.25, 1.15, 0.8], font_size=7.5)
    style.add_body(doc, "Las 14 familias tienen exportacion mecanica PASS para todos los casos listados. Solo 2 casos penales pasaron el subscore aislado; ninguno de los 214 permite FINAL. Longitud completa, calidad sustantiva, autoridades verificadas, aprobacion humana y calidad profesional no quedan demostradas.")

    add_section(doc, "5 Regresion, build y UI")
    style.add_body(doc, "La corrida actual `v5-global-current.json` paso 4,053 de 4,123 tests; 25 fallaron y 45 quedaron skipped. Las fallas vigentes se encuentran en `audit/generator-master/FAILURES.md`. En sintesis: cinco tests de acceptance intentan FINAL con dependencias pendientes; hay expectativas antiguas de marcadores, una expectativa de estado de postura, diferencias de cobertura/extraccion y fallos de materializacion/metadata. No se cambiaron gates para hacerlas verdes.")
    style.add_body(doc, "`npm run typecheck` paso con `tsc --noEmit`. `npm run build` no completo: Prisma reporto EPERM y al cierre PID 14312 seguia cargando `node_modules/.prisma/client/query_engine-windows.dll.node`. No se mato el proceso ni se borraron archivos. Una compilacion directa `npx next build --webpack` habia pasado antes, pero no reemplaza el build completo.")
    style.add_body(doc, "El test Playwright fue bloqueado antes de cargar la pagina por `listen EACCES: permission denied 127.0.0.1:3200`. Por ello page load, clicks, descarga DOCX/PDF y bloqueo FINAL por UI permanecen NOT RUN. Los 214 archivos del replay se crearon offline y no son descargas de interfaz.")

    add_section(doc, "6 Decisiones sobre repositorios externos")
    external = [
        ("Magic UI", "d7207e5", "MIT", "Referencia selectiva de componentes visuales; no incorporar como sistema completo."),
        ("shadcn/ui", "295a1f1", "MIT", "Referencia de componentes accesibles; conservar la arquitectura actual."),
        ("1Code", "9f1bc76", "Apache-2.0; archivado 2026-07-07", "No adoptar como base; solo patrones acotados tras revision."),
        ("it's free", "c53a15a", "No se encontro licencia en la raiz revisada", "No reutilizar codigo ni usar como proveedor; los free tiers son variables y no son apropiados para expedientes."),
    ]
    style.add_table(doc, ["Proyecto", "HEAD corto", "Licencia/estado", "Decision"], external,
                    [1.05, 0.82, 1.75, 2.7], font_size=7.2)
    style.add_body(doc, "No se clonaron repositorios, no se instalaron paquetes ni se copio codigo. HEADs completos y fuentes oficiales constan en `audit/external-repos/EXTERNAL_REPOS_DECISION.md`.")

    doc.add_page_break()
    add_section(doc, "7 Inventario de archivos del working tree",
                "Se conservaron cambios preexistentes. El inventario observado al cierre contiene 39 archivos versionados modificados y 22 nuevos/no versionados; la lista no atribuye todos a una unica fase.")
    add_path_bullets(doc, "Produccion modificada", [
        "app/api/legal-engine/appeal/reasoning-classification/route.ts",
        "app/machotes/components/CaseDocumentsReader.tsx",
        "app/machotes/page.tsx",
        "components/legal-taxonomy/LegalCatalogNavigator.tsx",
        "lib/ai/orchestrator.ts",
        "lib/ai/providers/gemini.ts",
        "lib/ai/providers/groq.ts",
        "lib/ai/providers/nvidia.ts",
        "lib/catalog/legalCatalog.ts",
        "lib/legal-engine/case-extraction/appealReasoningCandidates.ts",
        "lib/legal-engine/contestacionStructure.ts",
        "lib/legal-engine/documentRouting.ts",
        "lib/legal-engine/documentStrategies.ts",
        "lib/legal-engine/generatedLegalAdmission.ts",
        "lib/legal-engine/generationRequestContract.ts",
        "lib/legal-engine/pendingFields.ts",
        "lib/legal-engine/pipeline.ts",
        "lib/legal-engine/qualityGate.ts",
        "lib/legal-engine/seedMarkers.ts",
        "lib/operational-manual/core.ts",
    ])
    add_path_bullets(doc, "Pruebas versionadas modificadas", [
        "tests/ai/generationProviderTrace.test.ts; tests/ai/multiModelOrchestrator.test.ts",
        "tests/ai/nvidia404.test.ts; tests/ai/nvidiaOnly.test.ts; tests/ai/nvidiaPrimary.test.ts; tests/ai/nvidiaStructuredOutput.test.ts",
        "tests/audit/realAppealReasoningSource.test.ts",
        "tests/components/appealPhase1b.test.tsx; tests/components/appealReasoningCandidates.test.tsx; tests/components/appealResolutionReview.test.tsx",
        "tests/legal-engine/appealReasoningImpact.test.ts; tests/legal-engine/fallbackE2E.test.ts; tests/legal-engine/generationRequestContract.test.ts",
        "tests/legal-engine/loop8iCorporativoContractualPI.test.ts; tests/legal-engine/penalFallbackContract.test.ts",
        "tests/legal-taxonomy/legalCatalogRegistry.test.ts",
        "tests/operational-manual/importCanonical.test.ts; tests/operational-manual/manualCore.test.ts; tests/operational-manual/manualDraftingContextPurity.test.ts",
    ])
    add_path_bullets(doc, "Archivos nuevos del working tree", [
        "app/machotes/components/WritingAvailabilityNotice.tsx",
        "scripts/audit/build-writing-coverage.mjs; scripts/audit/create_generator_master_change_report.py; scripts/audit/create_generator_master_v4_closure.py",
        "scripts/audit/create_generator_master_v5_status_report.py",
        "scripts/audit/family-scorecard.ts; scripts/audit/update-v5-evidence.mjs; scripts/audit/v5-ui-draft-export.mjs",
        "tests/ai/providerErrorBodyPrivacy.test.ts; tests/api/appealReasoningClassificationAccess.test.ts",
        "tests/audit/familyScorecard.test.ts; tests/audit/generatorFamilyReplay.test.ts; tests/audit/generatorMasterInventory.test.ts",
        "tests/components/contestacionesCatalogCapabilities.test.tsx; tests/components/writingAvailabilityNotice.test.tsx",
        "tests/legal-engine/pendingMarkerNormalization.test.ts; tests/legal-engine/v5CivilPetitionContract.test.ts; tests/legal-engine/v5LegacyEvidenceSeed.test.ts; tests/legal-engine/v5PendingDependencyContract.test.ts",
        "tests/legal-taxonomy/contestacionesCapability.test.ts; tests/legal-taxonomy/functionalTypeStatus.test.ts; tests/legal-taxonomy/implementedDocumentContracts.test.ts",
    ])
    style.add_body(doc, "La lista amplia de evidencias de auditoria incluye `audit/generator-master/V5_EVIDENCE_SUMMARY.json`, `FAMILY_CLOSURE.md`, `GLOBAL_CLOSURE.md`, `FAILURES.md`, `STATUS.md`, `WRITING_COVERAGE_MATRIX.md`, `writing-coverage.json`, `v5-global-current.json`, el replay offline y `V5_CURRENT_CHANGES.patch`. El patch contiene el diff versionado; los archivos nuevos se enumeran arriba.")

    add_section(doc, "8 Bloqueos y siguiente criterio")
    for item in [
        "Resolver las 25 fallas actuales con diagnostico por causa y sin debilitar controles; actualizar solo expectativas demostrablemente obsoletas.",
        "Reintentar `npm run build` cuando el proceso propietario libere la DLL de Prisma; no borrar la DLL ni terminar PID 14312 automaticamente.",
        "Permitir localhost al runner para completar el E2E Playwright por clicks y verificar descargas DOCX/PDF y bloqueo FINAL desde UI.",
        "No usar provider real hasta contar con autorizacion y llaves configuradas localmente; no incorporar expediente privado a servicios gratuitos.",
        "Mantener los 277 tipos no certificados, la revision humana y FINAL bloqueados. No pasar a empaquetado Windows ni declarar GENERATOR_READY.",
    ]:
        style.add_bullet(doc, item)

    props = doc.core_properties
    props.title = "Estado V5 del generador juridico"
    props.subject = "LEX PLANTILLAS - cambios, pruebas y bloqueos del generador"
    props.author = ""
    props.last_modified_by = ""
    props.keywords = "LEX PLANTILLAS, generador, V5, auditoria, estado"
    doc.save(OUTPUT)
    print(f"DOCX_CREATED {OUTPUT}")


if __name__ == "__main__":
    build()
