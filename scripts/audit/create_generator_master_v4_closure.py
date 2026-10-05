from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
AUDIT = ROOT / "audit" / "generator-master"
SOURCE = AUDIT / "V4_CLOSURE.md"
OUTPUT = AUDIT / "GENERATOR_MASTER_V4_CLOSURE_REPORT.docx"
BLACK = RGBColor(0, 0, 0)
CODE = RGBColor(43, 73, 97)


def set_font(run, *, size=9.5, bold=None, color=BLACK, name="Arial"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def append_inline_markdown(paragraph, value: str):
    pieces = re.split(r"(`[^`]+`)", value)
    for piece in pieces:
        if not piece:
            continue
        if piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            set_font(run, size=8.5, color=CODE, name="Consolas")
        else:
            run = paragraph.add_run(piece)
            set_font(run)


def build():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for name, size in (("Title", 22), ("Heading 1", 14), ("Heading 2", 11)):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.keep_with_next = True

    content = SOURCE.read_text(encoding="utf-8").splitlines()
    for line in content:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.paragraph_format.space_after = Pt(7)
            run = paragraph.add_run(stripped[2:].strip())
            set_font(run, size=22, bold=True)
        elif stripped.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(5)
            run = paragraph.add_run(stripped[3:].strip())
            set_font(run, size=14, bold=True)
        elif stripped.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(stripped[4:].strip())
            set_font(run, size=11, bold=True)
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.24)
            paragraph.paragraph_format.first_line_indent = Inches(-0.14)
            paragraph.paragraph_format.space_after = Pt(3)
            append_inline_markdown(paragraph, stripped[2:].strip())
        else:
            paragraph = document.add_paragraph(style="Normal")
            append_inline_markdown(paragraph, stripped)

    document.core_properties.title = "Generator Master Closure V4"
    document.core_properties.subject = "LEX PLANTILLAS cambios A B C y verificación"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.keywords = "LEX PLANTILLAS, generador, cierre V4, auditoría"
    document.save(OUTPUT)
    print(f"DOCX_CREATED {OUTPUT}")


if __name__ == "__main__":
    build()
