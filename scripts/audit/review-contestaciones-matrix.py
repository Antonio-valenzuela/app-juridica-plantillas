"""Offline artifact inspection only. Never regenerates legal text or changes run artifacts."""
import hashlib
import json
import pathlib
import subprocess
import sys
from pypdf import PdfReader
from docx import Document
from PIL import Image, ImageOps, ImageDraw

root = pathlib.Path(sys.argv[1]).resolve()
output = root / 'offline-review'
output.mkdir(exist_ok=True)
results = []
for directory in sorted(root.iterdir()):
    if not (directory / 'result.json').exists():
        continue
    result = json.loads((directory / 'result.json').read_text(encoding='utf-8'))
    fixture = json.loads((directory / 'fixture.json').read_text(encoding='utf-8'))
    record = {'matter': fixture['matter'], 'generationResult': result.get('status'), 'exports': result.get('exports', {})}
    doc_path = directory / 'generated-document.json'
    if doc_path.exists():
        doc = json.loads(doc_path.read_text(encoding='utf-8'))
        text = '\n'.join(b['text'] for s in doc['sections'] for b in s['content'])
        identity_text = '\n'.join(b['text'] for s in doc['sections'] if s['type'] == 'identity' for b in s['content'])
        record.update({
            'assembledWords': len(text.split()), 'parties': doc.get('parties'),
            'identity': {'represented': fixture['represented'] in identity_text, 'capacity': fixture['capacity'] in identity_text, 'expediente': fixture['expediente'] in text, 'authority': fixture['authority'] in text},
            'sections': [{'id': s['id'], 'title': s['title'], 'words': sum(len(b['text'].split()) for b in s['content'])} for s in doc['sections']],
            'qualityGate': doc.get('qualityGate', doc.get('generationMetadata', {}).get('qualityGate')),
            'readiness': doc.get('generationMetadata', {}).get('readiness'),
            'validation': doc.get('validation'), 'coverage': doc.get('coverageMatrix'),
        })
        trace_path = directory / 'generation-trace.json'
        if trace_path.exists():
            trace = json.loads(trace_path.read_text(encoding='utf-8'))
            record['tasks'] = trace.get('taskExecutions', [])
            record['taskAccounting'] = trace.get('taskAccounting')
            record['issueAttempts'] = trace.get('issueGenerationAttempts', [])
            record['qualityGateTrace'] = trace.get('qualityGateResult')
            record['wordAccounting'] = trace.get('wordAccounting')
            record['providerModels'] = sorted(set(f"{t.get('providerActuallyUsed')}/{t.get('model')}" for t in trace.get('taskExecutions',[]) if t.get('model')))
            record['continuations'] = [{'taskId':t.get('taskId'),'count':t.get('continuationCount'),'status':t.get('responseStatus')} for t in trace.get('taskExecutions',[]) if t.get('continuationCount',0) or 'truncat' in str(t.get('responseStatus','')).lower()]
            coverage = trace.get('coverageMatrixAfterGeneration') or {}
            record['coverageTrace'] = coverage
        for format in ['pdf','docx']:
            path = directory / f'draft.{format}'
            if not path.exists():
                continue
            record[f'{format}Sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
            if format == 'docx':
                word = Document(path)
                exported = '\n'.join(p.text for p in word.paragraphs)
                record['docxParagraphs'] = len(word.paragraphs)
                record['docxWords'] = len(exported.split())
                (output / f'{directory.name}-docx-text.txt').write_text(exported, encoding='utf-8')
                record['docxVisual'] = 'NOT_EXECUTED: LibreOffice absent; OOXML parsed, not visually certified'
            else:
                pdf = PdfReader(path)
                record['pdfPages'] = len(pdf.pages)
                record['pdfWords'] = sum(len((p.extract_text() or '').split()) for p in pdf.pages)
                (output / f'{directory.name}-pdf-text.txt').write_text('\n'.join(p.extract_text() or '' for p in pdf.pages), encoding='utf-8')
                prefix = output / f'{directory.name}-page'
                subprocess.run(['pdftoppm','-scale-to','1000','-png',str(path),str(prefix)],check=True,capture_output=True)
                tiles = []
                for page in sorted(output.glob(f'{directory.name}-page-*.png')):
                    pic = Image.open(page).convert('RGB')
                    pic.thumbnail((660,850))
                    tile = Image.new('RGB',(680,880),'#eee8df')
                    tile.paste(pic,((680-pic.width)//2,22))
                    ImageDraw.Draw(tile).text((10,5),page.name,fill='black')
                    tiles.append(tile)
                sheet = Image.new('RGB',(680*min(3,len(tiles)),880*((len(tiles)+2)//3)),'white')
                for i,tile in enumerate(tiles): sheet.paste(tile,((i%3)*680,(i//3)*880))
                sheet.save(output / f'{directory.name}-contact.png')
    results.append(record)
(output / 'review.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps([{'matter':r['matter'],'identity':r.get('identity'),'words':r.get('assembledWords'),'pages':r.get('pdfPages'),'exports':r['exports']} for r in results],ensure_ascii=False,indent=2))
