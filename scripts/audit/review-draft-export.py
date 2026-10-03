"""Inspect downloaded HTTP artifacts, never modify the source run or call providers."""
import json
import pathlib
import subprocess
from docx import Document
from pypdf import PdfReader
from PIL import Image, ImageDraw

root = pathlib.Path('audit/draft-export-recovery/http').resolve()
doc = Document(root / 'current-editor-draft.docx')
word_text = '\n'.join(p.text for p in doc.paragraphs)
for section in doc.sections:
    word_text += '\n' + '\n'.join(p.text for p in section.header.paragraphs)
pdf = PdfReader(root / 'current-editor-draft.pdf')
pdf_text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
marker = 'EDICION ACTUAL DEL ABOGADO: copia de prueba para revision, no presentar.'
result = {'docxOpenedOOXML': True, 'docxVisual': 'NOT_EXECUTED: LibreOffice unavailable',
          'pdfPages': len(pdf.pages), 'docxParagraphs': len(doc.paragraphs), 'checks': {}}
for name, text in [('docx', word_text), ('pdf', pdf_text)]:
    checks = {'currentEditorText': marker in text, 'draftWarning': 'BORRADOR NO GUARDADO EN EL EXPEDIENTE' in text,
              'heading': 'PETITORIOS' in text, 'unverifiedCitationWarning': 'NO VERIFICADO' in text}
    result['checks'][name] = checks
    assert all(checks.values()), (name, checks)
    (root / f'{name}-text.txt').write_text(text, encoding='utf-8')
subprocess.run(['pdftoppm', '-scale-to', '1000', '-png', str(root / 'current-editor-draft.pdf'), str(root / 'pdf-page')], check=True, capture_output=True)
tiles = []
for page in sorted(root.glob('pdf-page-*.png')):
    pic = Image.open(page).convert('RGB')
    pic.thumbnail((660, 850))
    tile = Image.new('RGB', (680, 880), '#eee8df')
    tile.paste(pic, ((680-pic.width)//2, 22))
    ImageDraw.Draw(tile).text((10, 5), page.name, fill='black')
    tiles.append(tile)
sheet = Image.new('RGB', (680*min(3, len(tiles)), 880*((len(tiles)+2)//3)), 'white')
for i, tile in enumerate(tiles): sheet.paste(tile, ((i%3)*680, (i//3)*880))
sheet.save(root / 'pdf-contact.png')
(root / 'artifact-review.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
print(json.dumps(result, indent=2))
